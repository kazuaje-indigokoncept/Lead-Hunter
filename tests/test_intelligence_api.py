import csv
import io
import json
import sqlite3
import unittest
from unittest.mock import AsyncMock, patch

from app import db, intelligence, jobs
import test_maps_api


class IntelligenceApiTests(unittest.TestCase):
    # Reuse only the isolated database fixture, not the inherited test cases.
    setUp = test_maps_api.CampaignApiTests.setUp
    tearDown = test_maps_api.CampaignApiTests.tearDown
    def seed(self):
        self.body["options"] = {"sales_profile": "web", "maps_mode": "web"}
        cid = self.client.post("/api/campaigns", json=self.body).json()["id"]
        rows = [
            {"dedupe_key": "no-web", "name": "Diseño candidato", "source": "google_maps", "phone": "+573001234567", "reviews": 50},
            {"dedupe_key": "seo", "name": "SEO candidato", "source": "google_maps", "website": "https://negocio.test", "emails": ["hola@negocio.test"], "phone": "+573001234567", "reviews": 50,
             "web_audit": {"status": "ok", "has_meta_description": False, "has_title": False, "has_viewport": True, "has_contact_form": True, "https": True}},
        ]
        for row in rows:
            row["campaign_id"] = cid
            row["provenance"] = intelligence.initial_provenance(row, "google_maps_web")
            row["score"] = jobs.score(row)
            db.insert("leads", row)
        return cid

    def test_service_profile_changes_ranking_and_paging(self):
        cid = self.seed()
        web = self.client.get(f"/api/leads?campaign_id={cid}&sales_profile=web&sort=opportunity&per_page=1").json()
        seo = self.client.get(f"/api/leads?campaign_id={cid}&sales_profile=seo&sort=opportunity&per_page=1").json()
        self.assertEqual(web["items"][0]["name"], "Diseño candidato")
        self.assertEqual(seo["items"][0]["name"], "SEO candidato")
        self.assertEqual(web["total"], 2)

    def test_tag_opportunity_filters_match_export(self):
        cid = self.seed()
        query = f"campaign_id={cid}&tag=website:missing&opportunity=no_website&sales_profile=web&sort=opportunity"
        listing = self.client.get("/api/leads?" + query).json()
        self.assertEqual(listing["total"], 1)
        csv_result = self.client.get("/api/export?fmt=csv&" + query)
        rows = list(csv.reader(io.StringIO(csv_result.text.lstrip('\ufeff')), delimiter=";"))
        self.assertEqual(len(rows), 2)
        self.assertIn("Prioridad comercial", rows[0])
        self.assertIn("Procedencia", rows[0])
        self.assertEqual(rows[1][0], listing["items"][0]["name"])

    def test_excel_has_valid_new_columns(self):
        cid = self.seed()
        response = self.client.get(f"/api/export?campaign_id={cid}")
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(response.content))
        self.assertGreater(wb.active.max_column, 26)
        self.assertEqual(wb.active.max_row, 3)
        self.assertEqual(wb.active.cell(1, 24).value, "Prioridad comercial")
        wb.close()

    def test_recheck_merges_contacts_and_keeps_exact_source(self):
        cid = self.seed()
        db.set_setting("demo_mode", "0")
        lead = db.one("SELECT * FROM leads WHERE dedupe_key='seo'")
        data = {"emails": ["ventas@negocio.test"], "phones": ["+573111234567"], "linkedin": "https://www.linkedin.com/company/negocio",
            "web_audit": {"status": "ok", "pages": ["https://negocio.test/contacto"]},
            "provenance": {"emails": [intelligence.evidence("ventas@negocio.test", "website", "https://negocio.test/contacto")]}}
        with patch.object(jobs.website, "enrich", AsyncMock(return_value=data)):
            response = self.client.post(f"/api/leads/{lead['id']}/review-website?sales_profile=seo")
        self.assertEqual(response.status_code, 200)
        result = response.json()
        self.assertEqual(set(result["emails"]), {"hola@negocio.test", "ventas@negocio.test"})
        self.assertEqual(result["provenance"]["emails"][-1]["url"], "https://negocio.test/contacto")
        self.assertEqual(result["sales_profile"], "seo")

    def test_manual_links_persist_provenance_and_unchanged_save_keeps_audit(self):
        self.seed()
        lead = db.one("SELECT * FROM leads WHERE dedupe_key='no-web'")
        values = {"website": "https://negocio.test", "linkedin": "https://www.linkedin.com/company/negocio"}
        response = self.client.patch(f"/api/leads/{lead['id']}", json=values)
        self.assertEqual(response.status_code, 200)
        result = response.json()
        self.assertEqual(result["provenance"]["website"][-1]["source"], "manual")
        self.assertEqual(result["provenance"]["linkedin"][-1]["source"], "manual")
        db.update("leads", lead["id"], {"linkedin_data": {"status": "public_page"}})
        result = self.client.patch(f"/api/leads/{lead['id']}", json={**values, "notes": "Contacto pendiente"}).json()
        self.assertEqual(result["linkedin_data"]["status"], "public_page")
        self.assertEqual(self.client.patch(f"/api/leads/{lead['id']}", json={"linkedin": "https://linkedin.com.evil.test/company/a"}).status_code, 422)

    def test_export_formula_is_text_and_invalid_profile_rejected(self):
        self.seed()
        lead = db.one("SELECT * FROM leads WHERE dedupe_key='seo'")
        db.update("leads", lead["id"], {"name": "=1+1"})
        response = self.client.get("/api/export?fmt=csv")
        self.assertIn("'=1+1", response.text)
        self.assertEqual(self.client.get("/api/leads?sales_profile=invalid").status_code, 422)
        self.assertEqual(self.client.get("/api/leads?page=0").status_code, 422)

    def test_legacy_migration_preserves_records_and_is_idempotent(self):
        original = db.DB_PATH
        legacy = self.folder / "legacy.db"
        with sqlite3.connect(legacy) as conn:
            conn.executescript(db.SCHEMA)
            conn.execute("INSERT INTO campaigns(name,keywords,locations,sources) VALUES('Antigua','[]','[]','[]')")
            conn.execute("INSERT INTO leads(campaign_id,dedupe_key,source,name) VALUES(1,'legacy','rues','Negocio antiguo')")
        conn.close()
        try:
            db.DB_PATH = legacy
            db.init()
            db.init()
            lead = db.one("SELECT * FROM leads")
            self.assertEqual(lead["name"], "Negocio antiguo")
            self.assertEqual(lead["provenance"], {})
            self.assertEqual(lead["phones"], [])
            self.assertEqual(lead["web_audit"], {})
        finally:
            db.DB_PATH = original


class PublicNetworkTests(unittest.IsolatedAsyncioTestCase):
    async def test_private_and_local_targets_are_rejected(self):
        from app.enrich import website
        for url in ("http://127.0.0.1", "http://[::1]", "http://192.168.1.1", "file:///etc/passwd"):
            self.assertFalse(await website.public_url(url))
        async with __import__('httpx').AsyncClient(transport=__import__('httpx').MockTransport(lambda r: __import__('httpx').Response(200))) as client:
            data = await website.enrich(client, "http://127.0.0.1")
        self.assertEqual(data["web_audit"]["status"], "non_public_url")
