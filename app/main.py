"""LeadHunter - servidor web (FastAPI)."""
import asyncio
import csv
import io
import os
import json
from contextlib import asynccontextmanager
from pathlib import Path
from urllib.parse import urlparse

import httpx

from fastapi import Depends, FastAPI, HTTPException, Query, Request
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, field_validator

from . import auth, db, jobs, intelligence, companies
from .enrich import website, linkedin, instagram_public

STATIC = Path(__file__).resolve().parent.parent / "static"
SECURE_COOKIE = os.getenv("SECURE_COOKIE", "0") == "1"


@asynccontextmanager
async def lifespan(_app):
    # Reanuda lo interrumpido por un reinicio y arranca el programador de tandas.
    jobs.recover()
    task = asyncio.create_task(jobs.scheduler_loop())
    try:
        yield
    finally:
        task.cancel()


app = FastAPI(title="LeadHunter", docs_url=None, redoc_url=None, openapi_url=None, lifespan=lifespan)
app.mount("/static", StaticFiles(directory=STATIC), name="static")
db.init()
auth.init()
db.assign_orphans()
companies.backfill()

PUBLIC = {"/", "/api/login", "/api/logout", "/api/me"}


@app.middleware("http")
async def require_login(request: Request, call_next):
    request.state.user = auth.current_user(request)
    if request.url.path.startswith("/api/") and request.url.path not in PUBLIC \
            and not request.state.user:
        return JSONResponse({"detail": "Inicia sesión"}, status_code=401)
    resp = await call_next(request)
    resp.headers["X-Frame-Options"] = "DENY"
    resp.headers["X-Content-Type-Options"] = "nosniff"
    resp.headers["Referrer-Policy"] = "same-origin"
    return resp


# ---------------- sesión y usuarios ----------------
class LoginIn(BaseModel):
    username: str
    password: str


@app.post("/api/login")
def do_login(body: LoginIn, request: Request):
    token = auth.login(request, body.username, body.password)
    resp = JSONResponse({"ok": True})
    resp.set_cookie(auth.COOKIE, token, max_age=auth.SESSION_DAYS * 86400, httponly=True,
                    samesite="strict", secure=SECURE_COOKIE)
    return resp


@app.post("/api/logout")
def do_logout(request: Request):
    if request.cookies.get(auth.COOKIE):
        auth.logout(request.cookies[auth.COOKIE])
    resp = JSONResponse({"ok": True})
    resp.delete_cookie(auth.COOKIE)
    return resp


@app.get("/api/me")
def me(request: Request):
    if not request.state.user:
        raise HTTPException(401, "Inicia sesión")
    return request.state.user


class PasswordIn(BaseModel):
    current: str
    new: str


@app.post("/api/me/password")
def change_password(body: PasswordIn, request: Request):
    u = db.one("SELECT * FROM users WHERE id = ?", (request.state.user["id"],))
    if not auth.check_pw(body.current, u["pw_hash"]):
        raise HTTPException(400, "La contraseña actual no es correcta.")
    if len(body.new) < 8:
        raise HTTPException(400, "La nueva contraseña debe tener mín. 8 caracteres.")
    db.execute("UPDATE users SET pw_hash = ? WHERE id = ?", (auth.hash_pw(body.new), u["id"]))
    return {"ok": True}


class UserIn(BaseModel):
    username: str
    password: str
    is_admin: bool = False


@app.get("/api/users", dependencies=[Depends(auth.require_admin)])
def list_users():
    return db.query("""SELECT u.id, u.username, u.is_admin, u.active, u.created_at, u.last_login,
        (SELECT COUNT(*) FROM campaigns c WHERE c.owner_id = u.id) AS campaign_count,
        (SELECT COUNT(*) FROM campaigns c WHERE c.owner_id = u.id AND c.status IN ('programada','en curso')) AS active_campaigns,
        (SELECT COUNT(*) FROM leads l JOIN campaigns c ON c.id = l.campaign_id WHERE c.owner_id = u.id) AS lead_count,
        (SELECT COUNT(*) FROM sessions s WHERE s.user_id = u.id AND s.expires > CAST(strftime('%s','now') AS REAL)) AS open_sessions
        FROM users u ORDER BY u.is_admin DESC, u.username COLLATE NOCASE""")


class UserPatch(BaseModel):
    username: str | None = None
    is_admin: bool | None = None
    active: bool | None = None


@app.patch("/api/users/{uid}")
def update_user(uid: int, body: UserPatch, admin=Depends(auth.require_admin)):
    user = db.one("SELECT * FROM users WHERE id = ?", (uid,))
    if not user:
        raise HTTPException(404, "Usuario no encontrado")
    upd = {}
    if body.username is not None and body.username.strip() != user["username"]:
        name = body.username.strip()
        if len(name) < 3:
            raise HTTPException(400, "El usuario debe tener mín. 3 caracteres.")
        if db.one("SELECT id FROM users WHERE username = ? AND id != ?", (name, uid)):
            raise HTTPException(400, "Ese usuario ya existe.")
        upd["username"] = name
    for field in ("is_admin", "active"):
        value = getattr(body, field)
        if value is None or int(value) == user[field]:
            continue
        if uid == admin["id"]:
            raise HTTPException(400, "No puedes quitarte el rol de administrador ni desactivarte a ti mismo.")
        upd[field] = int(value)
    if not upd:
        return {"ok": True}
    if user["is_admin"] and user["active"] and (upd.get("is_admin") == 0 or upd.get("active") == 0):
        others = db.one("SELECT COUNT(*) n FROM users WHERE is_admin = 1 AND active = 1 AND id != ?", (uid,))["n"]
        if not others:
            raise HTTPException(400, "Debe quedar al menos un administrador activo.")
    db.update("users", uid, upd)
    if upd.get("active") == 0:
        db.execute("DELETE FROM sessions WHERE user_id = ?", (uid,))
    return {"ok": True}


@app.post("/api/users", dependencies=[Depends(auth.require_admin)])
def add_user(body: UserIn):
    return {"id": auth.create_user(body.username, body.password, body.is_admin)}


class ResetPasswordIn(BaseModel):
    password: str


@app.post("/api/users/{uid}/password", dependencies=[Depends(auth.require_admin)])
def reset_user_password(uid: int, body: ResetPasswordIn):
    if not db.one("SELECT id FROM users WHERE id = ?", (uid,)):
        raise HTTPException(404, "Usuario no encontrado")
    if len(body.password) < 8:
        raise HTTPException(400, "La contraseña debe tener mín. 8 caracteres.")
    db.execute("UPDATE users SET pw_hash = ? WHERE id = ?", (auth.hash_pw(body.password), uid))
    db.execute("DELETE FROM sessions WHERE user_id = ?", (uid,))
    return {"ok": True}


@app.delete("/api/users/{uid}")
def delete_user(uid: int, admin=Depends(auth.require_admin)):
    if uid == admin["id"]:
        raise HTTPException(400, "No puedes eliminar tu propio usuario.")
    # Sus campañas y empresas pasan al administrador que lo elimina; no se pierden datos.
    db.execute("UPDATE campaigns SET owner_id = ? WHERE owner_id = ?", (admin["id"], uid))
    db.execute("UPDATE companies SET owner_id = ? WHERE owner_id = ?", (admin["id"], uid))
    db.execute("DELETE FROM sessions WHERE user_id = ?", (uid,))
    db.execute("DELETE FROM users WHERE id = ?", (uid,))
    return {"ok": True}


# ---------------- acceso por usuario ----------------
def _user(request: Request):
    return request.state.user


def _owns(user, owner_id) -> bool:
    return bool(user["is_admin"]) or owner_id == user["id"]


def _owner_scope(user, column="owner_id"):
    """Condición SQL para limitar consultas a lo que el usuario puede ver."""
    return ("1=1", []) if user["is_admin"] else (f"{column} = ?", [user["id"]])

SECRET_KEYS = ["google_api_key", "meta_token", "ig_user_id", "socrata_token"]
PLAIN_KEYS = ["demo_mode", "price_per_1000"]


@app.get("/")
def index():
    return FileResponse(STATIC / "index.html")


# ---------------- ajustes ----------------
def _mask(v: str) -> str:
    return (v[:4] + "•" * 8 + v[-4:]) if len(v) > 10 else ("•" * len(v))


@app.get("/api/settings")
def get_settings():
    out = {k: _mask(db.get_setting(k)) for k in SECRET_KEYS}
    out["demo_mode"] = db.get_setting("demo_mode", "1")
    out["price_per_1000"] = db.get_setting("price_per_1000", "35")
    return out


@app.post("/api/settings", dependencies=[Depends(auth.require_admin)])
def save_settings(data: dict):
    for k, v in data.items():
        if k in SECRET_KEYS + PLAIN_KEYS and v is not None and "•" not in str(v):
            db.set_setting(k, str(v).strip())
    return get_settings()


# ---------------- campañas ----------------
class CampaignIn(BaseModel):
    name: str
    client: str = ""
    keywords: list[str]
    locations: list[str]
    sources: list[str] = Field(default_factory=lambda: ["google_maps", "website"])
    options: dict = Field(default_factory=dict)
    start: bool = True

    @field_validator("options")
    @classmethod
    def validate_options(cls, opts):
        if opts.get("instagram_mode", "api") not in ("api", "web"):
            raise ValueError("Modo de Instagram inválido")
        if opts.get("maps_mode", "api") not in ("api", "web"):
            raise ValueError("Modo de Google Maps inválido")
        if opts.get("sales_profile", "general") not in intelligence.PROFILES:
            raise ValueError("Perfil comercial inválido")
        web_pages = opts.get("web_max_pages", 4)
        if isinstance(web_pages, bool) or not isinstance(web_pages, int) or not 0 <= web_pages <= 8:
            raise ValueError("web_max_pages debe estar entre 0 y 8")
        for key, default, low, high in (("maps_max_results", 40, 1, 120), ("max_pages", 3, 1, 3)):
            value = opts.get(key, default)
            if isinstance(value, bool) or not isinstance(value, int) or not low <= value <= high:
                raise ValueError(f"{key} debe ser entero entre {low} y {high}")
        delay = opts.get("maps_delay", 2)
        if isinstance(delay, bool) or not isinstance(delay, (int, float)) or not 1 <= delay <= 30:
            raise ValueError("maps_delay debe estar entre 1 y 30 segundos")
        schedule = opts.get("schedule")
        if schedule is not None:
            opts["schedule"] = validate_schedule(schedule)
        rating = opts.get("min_rating")
        if rating is not None and (isinstance(rating, bool) or not isinstance(rating, (int, float)) or not 0 <= rating <= 5):
            raise ValueError("Rating mínimo debe estar entre 0 y 5")
        return opts


def validate_schedule(sch):
    if not isinstance(sch, dict):
        raise ValueError("Programación inválida")
    interval, per_batch = sch.get("interval_minutes", 60), sch.get("combos_per_batch", 1)
    for value in (interval, per_batch):
        if isinstance(value, bool) or not isinstance(value, int):
            raise ValueError("La programación usa números enteros")
    if not jobs.MIN_INTERVAL <= interval <= jobs.MAX_INTERVAL:
        raise ValueError(f"El intervalo debe estar entre {jobs.MIN_INTERVAL} minutos y 7 días")
    if not 1 <= per_batch <= jobs.MAX_COMBOS_PER_BATCH:
        raise ValueError(f"Las búsquedas por tanda deben estar entre 1 y {jobs.MAX_COMBOS_PER_BATCH}")
    return {"enabled": bool(sch.get("enabled")), "interval_minutes": interval, "combos_per_batch": per_batch}


def _clean(lst):
    return [x.strip() for x in lst if x and x.strip()]


@app.post("/api/estimate")
def estimate(c: CampaignIn):
    combos = len(_clean(c.keywords)) * len(_clean(c.locations))
    pages = int(c.options.get("max_pages", 3))
    web = c.options.get("maps_mode", "api") == "web"
    has_maps = "google_maps" in c.sources
    calls = combos * pages if has_maps and not web else 0
    price = float(db.get_setting("price_per_1000", "35"))
    return {"queries": combos, "max_calls": calls,
            "max_results": combos * c.options.get("maps_max_results", 40) if has_maps and web else calls * 20,
            "maps_mode": "web" if web else "api",
            "max_cost_usd": round(calls * price / 1000, 2)}


def _schedule_info(c):
    sch = jobs.schedule_of(c)
    total = len(c["keywords"]) * len(c["locations"])
    return {"schedule": sch, "total_combos": total, "batch_cursor": c.get("batch_cursor") or 0,
            "next_run_at": c.get("next_run_at"),
            "batches_total": -(-total // sch["combos_per_batch"]) if sch["enabled"] else 1}


@app.get("/api/campaigns")
def list_campaigns(request: Request, owner_id: int | None = None):
    user = _user(request)
    scope, params = _owner_scope(user, "c.owner_id")
    if owner_id and user["is_admin"]:
        scope, params = "c.owner_id = ?", [owner_id]
    rows = db.query(f"""SELECT c.id, c.name, c.client, c.status, c.progress, c.api_calls,
        c.created_at, c.keywords, c.locations, c.sources, c.owner_id, c.options,
        c.next_run_at, c.batch_cursor,
        (SELECT username FROM users u WHERE u.id = c.owner_id) AS owner_name,
        (SELECT COUNT(*) FROM leads l WHERE l.campaign_id = c.id) AS total,
        (SELECT COUNT(*) FROM leads l WHERE l.campaign_id = c.id AND l.emails != '[]') AS with_email
        FROM campaigns c WHERE {scope} ORDER BY c.id DESC""", tuple(params))
    for r in rows:
        r["running"] = r["id"] in jobs.RUNNING
        r.update(_schedule_info(r))
        r.pop("options", None)
    return rows


@app.post("/api/campaigns")
async def create_campaign(c: CampaignIn, request: Request):
    kws, locs = _clean(c.keywords), _clean(c.locations)
    if not kws or not locs:
        raise HTTPException(400, "Agrega al menos una palabra clave y una ubicación.")
    cid = db.insert("campaigns", {"name": c.name.strip() or f"{kws[0]} · {locs[0]}",
                                  "client": c.client.strip(), "keywords": kws,
                                  "locations": locs, "sources": c.sources,
                                  "options": c.options, "owner_id": _user(request)["id"]})
    if c.start:
        jobs.start(cid)
    return {"id": cid}


def _campaign(cid, user):
    c = db.one("SELECT * FROM campaigns WHERE id = ?", (cid,))
    # 404 también cuando es de otro usuario, para no revelar que existe.
    if not c or not _owns(user, c["owner_id"]):
        raise HTTPException(404, "Campaña no encontrada")
    return c


@app.get("/api/campaigns/{cid}")
def get_campaign(cid: int, request: Request):
    c = _campaign(cid, _user(request))
    c["running"] = cid in jobs.RUNNING
    owner = db.one("SELECT username FROM users WHERE id = ?", (c["owner_id"],))
    c["owner_name"] = owner["username"] if owner else ""
    c.update(_schedule_info(c))
    s = db.one("""SELECT COUNT(*) total,
        SUM(emails != '[]') email, SUM(whatsapp IS NOT NULL AND whatsapp != '') wa,
        SUM(phone IS NOT NULL AND phone != '') phone,
        SUM(website IS NULL OR website = '') no_web,
        SUM(instagram IS NOT NULL AND instagram != '') ig,
        ROUND(AVG(score),0) avg_score
        FROM leads WHERE campaign_id = ?""", (cid,))
    c["stats"] = {k: (v or 0) for k, v in s.items()}
    return c


@app.post("/api/campaigns/{cid}/start")
async def start_campaign(cid: int, request: Request):
    _campaign(cid, _user(request))
    return {"started": jobs.start(cid)}


@app.post("/api/campaigns/{cid}/stop")
def stop_campaign(cid: int, request: Request):
    _campaign(cid, _user(request))
    return {"stopped": jobs.stop(cid)}


@app.delete("/api/campaigns/{cid}")
def delete_campaign(cid: int, request: Request):
    _campaign(cid, _user(request))
    jobs.stop(cid)
    db.execute("DELETE FROM leads WHERE campaign_id = ?", (cid,))
    db.execute("DELETE FROM campaigns WHERE id = ?", (cid,))
    return {"ok": True}


class ScheduleIn(BaseModel):
    enabled: bool
    interval_minutes: int = 60
    combos_per_batch: int = 1


@app.put("/api/campaigns/{cid}/schedule")
def set_schedule(cid: int, body: ScheduleIn, request: Request):
    c = _campaign(cid, _user(request))
    try:
        sch = validate_schedule(body.model_dump())
    except ValueError as e:
        raise HTTPException(422, str(e))
    db.update("campaigns", cid, {"options": {**c["options"], "schedule": sch}})
    if cid not in jobs.RUNNING:
        if not sch["enabled"] and c["status"] in ("programada", "pausada"):
            db.update("campaigns", cid, {"status": "detenida", "next_run_at": None})
        elif sch["enabled"] and c["status"] == "programada" and c.get("next_run_at"):
            # Un intervalo más corto adelanta la próxima tanda; uno más largo la retrasa.
            last = c["next_run_at"] - jobs.schedule_of(c)["interval_minutes"] * 60
            db.update("campaigns", cid, {"next_run_at": max(last + sch["interval_minutes"] * 60, 0)})
    return get_campaign(cid, request)


# ---------------- leads ----------------
SORTS = {"score": "score DESC", "rating": "rating DESC", "reviews": "reviews DESC",
         "name": "name COLLATE NOCASE ASC", "recent": "id DESC"}


def decorate(lead, profile=""):
    if not profile:
        c = db.one("SELECT options FROM campaigns WHERE id = ?", (lead["campaign_id"],))
        profile = (c or {}).get("options", {}).get("sales_profile", "general")
    cid = companies.attach(lead)
    company = db.one("SELECT name FROM companies WHERE id=?", (cid,))
    return {**lead, "company_id": cid, "company_name": company["name"], **intelligence.analyze(lead, profile)}


def analyzed_rows(where, params, sales_profile="", opportunity="", tag="", min_priority=0, sort="score"):
    if sales_profile and sales_profile not in intelligence.PROFILES:
        raise HTTPException(422, "Perfil comercial inválido")
    campaigns = {c["id"]: c["options"].get("sales_profile", "general") for c in db.query("SELECT id, options FROM campaigns")}
    rows = [decorate(r, sales_profile or campaigns.get(r["campaign_id"], "general"))
            for r in db.query(f"SELECT * FROM leads WHERE {where} ORDER BY {SORTS.get(sort, 'score DESC')}", tuple(params))]
    rows = [r for r in rows if r["opportunity_score"] >= min_priority
            and (not opportunity or any(o["key"] == opportunity for o in r["opportunities"]))
            and (not tag or any(f"{t['key']}:{t['state']}" == tag for t in r["tags"]))]
    if sort == "opportunity":
        rows.sort(key=lambda r: (r["opportunity_score"], r["score"], r["id"]), reverse=True)
    return rows


@app.get("/api/sales-profiles")
def sales_profiles():
    return [{"id": key, "label": label} for key, label in intelligence.PROFILES.items()]


def _lead_filter(user, campaign_id, q, has_email, has_wa, no_web, has_ig, min_score, status, source):
    where, p = ["1=1"], []
    if not user["is_admin"]:
        where.append("campaign_id IN (SELECT id FROM campaigns WHERE owner_id = ?)"); p.append(user["id"])
    if campaign_id:
        where.append("campaign_id = ?"); p.append(campaign_id)
    if q:
        where.append("(name LIKE ? OR address LIKE ? OR category LIKE ? OR emails LIKE ?)")
        p += [f"%{q}%"] * 4
    if has_email:
        where.append("emails != '[]'")
    if has_wa:
        where.append("whatsapp IS NOT NULL AND whatsapp != ''")
    if no_web:
        where.append("(website IS NULL OR website = '')")
    if has_ig:
        where.append("instagram IS NOT NULL AND instagram != ''")
    if min_score:
        where.append("score >= ?"); p.append(min_score)
    if status:
        where.append("status = ?"); p.append(status)
    if source:
        where.append("source LIKE ?"); p.append(f"%{source}%")
    return " AND ".join(where), p


@app.get("/api/leads")
def list_leads(request: Request, campaign_id: int | None = None, q: str = "", has_email: bool = False,
               has_wa: bool = False, no_web: bool = False, has_ig: bool = False,
               min_score: int = 0, status: str = "", source: str = "",
               sort: str = "score", page: int = Query(1, ge=1), per_page: int = Query(50, ge=1, le=500),
               sales_profile: str = "", opportunity: str = "", tag: str = "", min_priority: int = Query(0, ge=0, le=100)):
    where, p = _lead_filter(_user(request), campaign_id, q, has_email, has_wa, no_web, has_ig,
                            min_score, status, source)
    rows = analyzed_rows(where, p, sales_profile, opportunity, tag, min_priority, sort)
    return {"total": len(rows), "page": page, "per_page": per_page,
            "items": rows[(page - 1) * per_page:page * per_page]}


class LeadPatch(BaseModel):
    status: str | None = None
    notes: str | None = None
    linkedin: str | None = None
    website: str | None = None
    instagram: str | None = None

    @field_validator("instagram")
    @classmethod
    def instagram_url(cls, value):
        if not value:
            return value
        try:
            valid = instagram_public.profile_url(value)
        except ValueError:
            valid = None
        if not valid:
            raise ValueError("Introduce una URL HTTPS de perfil de Instagram")
        return valid

    @field_validator("website")
    @classmethod
    def website_url(cls, value):
        if not value:
            return value
        parsed = urlparse(value)
        if parsed.scheme not in ("http", "https") or not parsed.hostname or parsed.username:
            raise ValueError("Introduce una URL HTTP o HTTPS de la web del negocio")
        return value

    @field_validator("linkedin")
    @classmethod
    def linkedin_url(cls, value):
        if not value:
            return value
        parsed = urlparse(value)
        host = parsed.hostname or ""
        if parsed.scheme != "https" or parsed.username or not (host == "linkedin.com" or host.endswith(".linkedin.com")) or not parsed.path.startswith(("/company/", "/in/")):
            raise ValueError("Introduce una URL HTTPS de empresa o perfil de LinkedIn")
        return value


def _lead(lid, user):
    lead = db.one("SELECT * FROM leads WHERE id = ?", (lid,))
    if not lead or not _owns(user, companies.lead_owner(lead)):
        raise HTTPException(404, "Lead no encontrado")
    return lead


@app.patch("/api/leads/{lid}")
def patch_lead(lid: int, body: LeadPatch, request: Request):
    lead = _lead(lid, _user(request))
    upd = body.model_dump(exclude_none=True)
    if body.website is not None and body.website != (lead.get("website") or ""):
        added = {"website": [intelligence.evidence(body.website, "manual", body.website, verification="aportado", confidence="baja")]} if body.website else {}
        upd["provenance"] = intelligence.merge_provenance(lead.get("provenance"), added)
        upd.update(web_audit={}, enriched=0)
    if body.linkedin is not None and body.linkedin != (lead.get("linkedin") or ""):
        added = {"linkedin": [intelligence.evidence(body.linkedin, "manual", body.linkedin, verification="aportado", confidence="baja")]} if body.linkedin else {}
        upd["provenance"] = intelligence.merge_provenance(upd.get("provenance", lead.get("provenance")), added)
        upd["linkedin_data"] = {}
    if body.instagram is not None and body.instagram != (lead.get("instagram") or ""):
        added = {"instagram": [intelligence.evidence(body.instagram, "manual", body.instagram, verification="aportado", confidence="baja")]} if body.instagram else {}
        upd["provenance"] = intelligence.merge_provenance(upd.get("provenance", lead.get("provenance")), added)
        for field in ("ig_bio", "ig_followers"):
            upd["provenance"].pop(field, None)
        upd.update(instagram_data={}, ig_bio=None, ig_followers=None)
    upd["score"] = jobs.score({**lead, **upd})
    db.update("leads", lid, upd)
    return decorate(db.one("SELECT * FROM leads WHERE id = ?", (lid,)))


@app.get("/api/companies")
def list_companies(request: Request, campaign_id: int | None = None, q: str = "", sales_profile: str = "general",
                   page: int = Query(1, ge=1), per_page: int = Query(50, ge=1, le=200)):
    if sales_profile not in intelligence.PROFILES:
        raise HTTPException(422, "Perfil comercial inválido")
    companies.backfill()
    rows = []
    scope, params = _owner_scope(_user(request))
    for company in db.query(f"SELECT * FROM companies WHERE {scope} ORDER BY id DESC", tuple(params)):
        info = companies.detail(company["id"], sales_profile)
        records = info["records"]
        if not records or campaign_id and not any(r["campaign_id"] == campaign_id for r in records):
            continue
        if q and q.lower() not in " ".join(info["names"] + [r.get("city") or "" for r in records]).lower():
            continue
        emails = list(dict.fromkeys(e for r in records for e in r.get("emails", [])))
        phones = list(dict.fromkeys(p for r in records for p in ((r.get("phones") or []) + ([r["phone"]] if r.get("phone") else []))))
        sites = list(dict.fromkeys(r["website"] for r in records if r.get("website")))
        contacts = [c for c in info["contacts"] if c["relationship"] != "rejected"]
        analyses = [intelligence.analyze(r, sales_profile) for r in records]
        tags = []
        for field in intelligence.CONTACTS:
            candidates = [next(t for t in a["tags"] if t["key"] == field) for a in analyses]
            state = next((s for s in ("present", "inferred") if any(t["state"] == s for t in candidates)),
                         "missing" if all(t["state"] == "missing" for t in candidates) else "unknown")
            tag = next(t for t in candidates if t["state"] == state)
            tags.append({**tag, "detail": tag["detail"] + " Resumen de los registros asociados a esta empresa."})
        rows.append({"id": company["id"], "name": company["name"], "record_count": len(records),
            "campaign_count": len({r["campaign_id"] for r in records}), "emails": emails, "phones": phones, "websites": sites,
            "cities": list(dict.fromkeys(r["city"] for r in records if r.get("city"))),
            "sources": list(dict.fromkeys(r["source"] for r in records)), "contacts": contacts,
            "opportunity_score": max(intelligence.analyze(r, sales_profile)["opportunity_score"] for r in records),
            "sales_profile_label": intelligence.PROFILES[sales_profile], "tags": tags,
            "demo": any(a["data_quality"]["demo"] for a in analyses)})
    rows.sort(key=lambda r: (r["opportunity_score"], r["id"]), reverse=True)
    return {"items": rows[(page-1)*per_page:page*per_page], "total": len(rows), "page": page, "per_page": per_page}


def company_exists(cid, user):
    company = db.one("SELECT * FROM companies WHERE id=?", (cid,))
    if not company or not _owns(user, company["owner_id"]):
        raise HTTPException(404, "Empresa no encontrada")
    return company


@app.get("/api/companies/{cid}")
def get_company(cid: int, request: Request, sales_profile: str = "general"):
    company_exists(cid, _user(request))
    if sales_profile not in intelligence.PROFILES:
        raise HTTPException(422, "Perfil comercial inválido")
    info = companies.detail(cid, sales_profile)
    info["records"] = [decorate(r, sales_profile) for r in info["records"]]
    return info


class CompanyName(BaseModel):
    name: str = Field(min_length=1, max_length=200)


@app.patch("/api/companies/{cid}")
def rename_company(cid: int, body: CompanyName, request: Request):
    company_exists(cid, _user(request))
    if not body.name.strip():
        raise HTTPException(400, "Introduce el nombre de la empresa")
    db.update("companies", cid, {"name": body.name.strip()})
    return {"ok": True}


class AssignCompany(BaseModel):
    company_id: int | None = Field(None, ge=1)


@app.post("/api/leads/{lid}/company")
def assign_company(lid: int, body: AssignCompany, request: Request):
    user = _user(request)
    lead = existing_lead(lid, user)
    owner = companies.lead_owner(lead)
    if body.company_id:
        company = company_exists(body.company_id, user)
        # Un registro solo puede ir a una empresa del mismo dueño que su campaña.
        if company["owner_id"] != owner:
            raise HTTPException(404, "Empresa no encontrada")
        cid = body.company_id
    else:
        cid = db.insert("companies", {"name": lead.get("name") or "Empresa sin nombre", "owner_id": owner,
                                      "grouping_note": "Empresa separada manualmente."})
    db.update("leads", lid, {"company_id": cid, "company_locked": 1})
    return decorate(db.one("SELECT * FROM leads WHERE id=?", (lid,)))


class ContactIn(BaseModel):
    name: str = Field("", max_length=200)
    role: str = Field("", max_length=300)
    linkedin: str = ""
    emails: list[str] = Field(default_factory=list, max_length=5)
    phones: list[str] = Field(default_factory=list, max_length=5)
    source_url: str = ""
    relationship: str = "pending"

    @field_validator("linkedin")
    @classmethod
    def profile_url(cls, value):
        if value and not companies.linkedin_url(value, "in"):
            raise ValueError("Introduce un perfil HTTPS /in/ de LinkedIn")
        return companies.linkedin_url(value, "in")

    @field_validator("relationship")
    @classmethod
    def relation(cls, value):
        if value not in ("pending", "confirmed", "rejected"):
            raise ValueError("Estado de vínculo inválido")
        return value

    @field_validator("source_url")
    @classmethod
    def source(cls, value):
        if value and (urlparse(value).scheme not in ("http", "https") or not urlparse(value).hostname or urlparse(value).username):
            raise ValueError("Introduce una URL HTTP/HTTPS de evidencia")
        return value

    @field_validator("emails")
    @classmethod
    def email_values(cls, values):
        if any(not website.EMAIL_RE.fullmatch(v) for v in values):
            raise ValueError("Formato de correo inválido")
        return [v.lower() for v in values]

    @field_validator("phones")
    @classmethod
    def phone_values(cls, values):
        result = [website.clean_phone(v) for v in values]
        if any(v is None for v in result):
            raise ValueError("Formato de teléfono inválido")
        return result


@app.post("/api/companies/{cid}/contacts")
def add_company_contact(cid: int, body: ContactIn, request: Request):
    company_exists(cid, _user(request))
    if not body.name.strip() and not body.linkedin:
        raise HTTPException(400, "Añade el nombre o el perfil de LinkedIn")
    data = body.model_dump()
    demo = db.get_setting("demo_mode", "1") == "1" or any(intelligence.analyze(r)["data_quality"]["demo"] for r in companies.context(cid)["records"])
    data["evidence_records"] = [intelligence.evidence(
        f"{body.name} · {body.role} · vínculo {body.relationship}", "demo" if demo else "manual", body.source_url or body.linkedin,
        verification="demo" if demo else "aportado", confidence="media" if body.relationship == "confirmed" else "baja")]
    contact_id = companies.save_contact(cid, data, manual=True)
    return db.one("SELECT * FROM company_contacts WHERE id=?", (contact_id,))


class ContactDecision(BaseModel):
    relationship: str
    @field_validator("relationship")
    @classmethod
    def relation(cls, value):
        if value not in ("pending", "confirmed", "rejected"):
            raise ValueError("Estado de vínculo inválido")
        return value


@app.patch("/api/company-contacts/{contact_id}")
def decide_contact(contact_id: int, body: ContactDecision, request: Request):
    contact = db.one("""SELECT cc.*, co.owner_id FROM company_contacts cc
                        JOIN companies co ON co.id = cc.company_id WHERE cc.id=?""", (contact_id,))
    if not contact or not _owns(_user(request), contact["owner_id"]):
        raise HTTPException(404, "Contacto no encontrado")
    record = intelligence.evidence(f"Vínculo {body.relationship} por revisión manual", "manual", contact["linkedin"], verification="aportado", confidence="media")
    records = intelligence.merge_provenance({"r": contact["evidence_records"]}, {"r": [record]})["r"]
    db.update("company_contacts", contact_id, {"relationship": body.relationship, "evidence_records": records})
    return {"ok": True}


@app.post("/api/companies/{cid}/review-decision-makers")
async def review_decision_makers(cid: int, request: Request, sales_profile: str = "general"):
    company_exists(cid, _user(request))
    if sales_profile not in intelligence.PROFILES:
        raise HTTPException(422, "Perfil comercial inválido")
    info = companies.context(cid)
    if any(r["campaign_id"] in jobs.RUNNING for r in info["records"]):
        raise HTTPException(409, "Espera a que termine la campaña antes de revisar responsables")
    demo = db.get_setting("demo_mode", "1") == "1" or any(intelligence.analyze(r)["data_quality"]["demo"] for r in info["records"])
    async with httpx.AsyncClient(headers={"User-Agent": website.UA}) as client:
        seen = set()
        for lead in info["records"]:
            if lead.get("website") and lead["website"] not in seen:
                await jobs.enrich_lead(client, lead, demo, 8)
                seen.add(lead["website"])
            if len(seen) >= 5:
                break
        result = await linkedin.review_company_people(client, cid, demo=demo)
    return {"company": get_company(cid, request, sales_profile), "review": result}


def existing_lead(lid, user):
    lead = _lead(lid, user)
    if lead["campaign_id"] in jobs.RUNNING:
        raise HTTPException(409, "Espera a que termine la campaña para volver a revisar este lead")
    return lead


@app.post("/api/leads/{lid}/review-website")
async def review_website(lid: int, request: Request, sales_profile: str = ""):
    lead = existing_lead(lid, _user(request))
    if not lead.get("website"):
        raise HTTPException(400, "Este lead no tiene web para revisar")
    demo = db.get_setting("demo_mode", "1") == "1" or intelligence.analyze(lead)["data_quality"]["demo"]
    async with httpx.AsyncClient(headers={"User-Agent": website.UA}) as client:
        return decorate(await jobs.enrich_lead(client, lead, demo), sales_profile)


@app.post("/api/leads/{lid}/review-linkedin")
async def review_linkedin(lid: int, request: Request, sales_profile: str = ""):
    lead = existing_lead(lid, _user(request))
    if not lead.get("linkedin"):
        raise HTTPException(400, "Añade primero un enlace de LinkedIn")
    demo = db.get_setting("demo_mode", "1") == "1" or intelligence.analyze(lead)["data_quality"]["demo"]
    async with httpx.AsyncClient(headers={"User-Agent": website.UA}) as client:
        data = {"linkedin_data": {"status": "demo", "url": lead["linkedin"], "checked_at": intelligence.now()}} if demo else await linkedin.lookup(client, lead["linkedin"])
    return decorate(jobs.apply_enrichment(lead, data, demo=demo), sales_profile)


@app.post("/api/leads/{lid}/review-instagram")
async def review_instagram(lid: int, request: Request, sales_profile: str = "", mode: str = "web"):
    lead = existing_lead(lid, _user(request))
    if mode not in ("web", "api"):
        raise HTTPException(422, "Modo de Instagram inválido")
    if not lead.get("instagram"):
        raise HTTPException(400, "Añade primero un enlace de Instagram y guarda el lead")
    demo = db.get_setting("demo_mode", "1") == "1" or intelligence.analyze(lead)["data_quality"]["demo"]
    token, ig_id = db.get_setting("meta_token"), db.get_setting("ig_user_id")
    if mode == "api" and not demo and not (token and ig_id):
        raise HTTPException(400, "Faltan token de Meta / ID de cuenta IG en Ajustes")
    async with httpx.AsyncClient(headers={"User-Agent": website.UA}) as client:
        result = await jobs.enrich_instagram(client, lead, mode=mode, demo=demo, token=token, ig_id=ig_id)
    return decorate(result, sales_profile)


EXPORT_COLS = [("name", "Nombre"), ("category", "Categoría"), ("city", "Ciudad"),
               ("address", "Dirección"), ("phone", "Teléfono"), ("whatsapp", "WhatsApp"),
               ("emails", "Correos"), ("website", "Sitio web"), ("instagram", "Instagram"),
               ("ig_followers", "Seguidores IG"), ("facebook", "Facebook"),
               ("tiktok", "TikTok"), ("linkedin", "LinkedIn"), ("rating", "Rating"),
               ("reviews", "Reseñas"), ("score", "Score"), ("status", "Estado"),
               ("source", "Fuente"), ("maps_url", "Google Maps"), ("notes", "Notas"),
               ("extra", "Datos RUES"), ("phones", "Todos los teléfonos"),
               ("sales_profile_label", "Servicio"), ("opportunity_score", "Prioridad comercial"),
               ("opportunities", "Oportunidades"), ("tags", "Etiquetas"),
               ("data_quality", "Calidad y trazabilidad"), ("provenance", "Procedencia"),
               ("web_audit", "Revisión web"), ("linkedin_data", "Revisión LinkedIn"),
               ("ig_bio", "Biografía Instagram"), ("instagram_data", "Revisión Instagram")]
EXPORT_COLS += [("company_id", "ID empresa"), ("company_name", "Empresa agrupada")]


def _flat(v):
    if isinstance(v, list):
        if v and isinstance(v[0], dict):
            return ", ".join(x.get("title", x.get("label", str(x))) for x in v)
        return ", ".join(map(str, v))
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False) if v else ""
    return "" if v is None else v


@app.get("/api/export")
def export(request: Request, campaign_id: int | None = None, fmt: str = "xlsx", q: str = "",
           has_email: bool = False, has_wa: bool = False, no_web: bool = False,
           has_ig: bool = False, min_score: int = 0, status: str = "", source: str = "",
           sales_profile: str = "", opportunity: str = "", tag: str = "", min_priority: int = Query(0, ge=0, le=100), sort: str = "score"):
    user = _user(request)
    if campaign_id:
        _campaign(campaign_id, user)
    where, p = _lead_filter(user, campaign_id, q, has_email, has_wa, no_web, has_ig,
                            min_score, status, source)
    rows = analyzed_rows(where, p, sales_profile, opportunity, tag, min_priority, sort)
    name = "leads"
    if campaign_id:
        c = _campaign(campaign_id, user)
        name = "leads_" + "".join(ch if ch.isalnum() else "_" for ch in c["name"])[:40]
    header = [h for _, h in EXPORT_COLS]
    data = [[_flat(r.get(k)) for k, _ in EXPORT_COLS] for r in rows]
    data = [["'" + value if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")) else value for value in row] for row in data]

    if fmt == "csv":
        buf = io.StringIO()
        buf.write("﻿")  # BOM para que Excel lea bien las tildes
        w = csv.writer(buf, delimiter=";")
        w.writerow(header); w.writerows(data)
        return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv",
                                 headers={"Content-Disposition": f'attachment; filename="{name}.csv"'})

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = "Leads"
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
    for row in data:
        ws.append(row)
    widths = [30, 18, 14, 34, 16, 28, 34, 30, 30, 12, 30, 22, 22, 8, 9, 7, 12, 12, 36, 24, 40]
    for i, wdt in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = wdt
    for i in range(len(widths), len(EXPORT_COLS)):
        ws.column_dimensions[get_column_letter(i + 1)].width = 36
    ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO(); wb.save(out); out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'})
